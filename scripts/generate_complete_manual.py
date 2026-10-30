# -*- coding: utf-8 -*-
"""
COMMERCIA — Cahier de Fonctionnement Intégral & Manuel d'Utilisation Exhaustif
Générateur automatisé et ultra-complet du document Word (.docx) et PDF (.pdf)
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import inch, cm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

# =========================================================================
# 1. GÉNÉRATION WORD (.DOCX)
# =========================================================================

def build_docx_manual(output_path):
    doc = Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Arial'
    font_normal.size = Pt(10)
    font_normal.color.rgb = RGBColor(30, 41, 59)
    
    PRIMARY = "059669"
    PRIMARY_RGB = RGBColor(5, 150, 105)
    SECONDARY_RGB = RGBColor(30, 41, 59)
    ACCENT_RGB = RGBColor(37, 99, 235)
    MUTED_RGB = RGBColor(100, 116, 139)
    
    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def add_callout(text, title="EXEMPLE CONCRET & CAS PRATIQUE", bg_hex="F0FDF4", border_hex=PRIMARY):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.8)
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        run_title = p.add_run(f"📌 {title}\n")
        run_title.bold = True
        run_title.font.name = "Arial"
        run_title.font.size = Pt(10.5)
        run_title.font.color.rgb = RGBColor.from_string(border_hex)
        
        run_text = p.add_run(text)
        run_text.font.name = "Arial"
        run_text.font.size = Pt(9.5)
        run_text.font.color.rgb = RGBColor(51, 65, 85)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(24)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(17)
        run.font.color.rgb = PRIMARY_RGB
        p_line = doc.add_paragraph()
        p_line.paragraph_format.space_before = Pt(0)
        p_line.paragraph_format.space_after = Pt(12)
        run_line = p_line.add_run("―" * 55)
        run_line.font.color.rgb = PRIMARY_RGB
        run_line.font.size = Pt(11)

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(13)
        run.font.color.rgb = SECONDARY_RGB

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(11)
        run.font.color.rgb = ACCENT_RGB

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.bold = True
            run_b.font.color.rgb = SECONDARY_RGB
        run_t = p.add_run(text)
        run_t.font.color.rgb = SECONDARY_RGB
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.bold = True
            run_b.font.color.rgb = SECONDARY_RGB
        run_t = p.add_run(text)
        run_t.font.color.rgb = SECONDARY_RGB
        return p

    def add_styled_table(headers, rows, col_widths=None):
        tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        
        hdr_cells = tbl.rows[0].cells
        for i, header_text in enumerate(headers):
            cell = hdr_cells[i]
            if col_widths and i < len(col_widths):
                cell.width = Inches(col_widths[i])
            set_cell_background(cell, PRIMARY)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(header_text)
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(255, 255, 255)
            
        for r_idx, row in enumerate(rows):
            row_cells = tbl.rows[r_idx + 1].cells
            bg_col = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
            for c_idx, val in enumerate(row):
                cell = row_cells[c_idx]
                if col_widths and c_idx < len(col_widths):
                    cell.width = Inches(col_widths[c_idx])
                set_cell_background(cell, bg_col)
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(str(val))
                run.font.name = "Arial"
                run.font.size = Pt(9)
                run.font.color.rgb = SECONDARY_RGB
                
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # PAGE DE COUVERTURE
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(60)
    run_pre = p_pre.add_run("RÉPUBLIQUE DU COMMERCE & ENTREPRENEURIAT MODERNE\nMANUEL OFFICIEL D'EXPLOITATION OPÉRATIONNELLE")
    run_pre.font.size = Pt(10)
    run_pre.font.color.rgb = MUTED_RGB
    run_pre.bold = True

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(40)
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run("COMMERCIA")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(36)
    run_title.bold = True
    run_title.font.color.rgb = PRIMARY_RGB

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(24)
    run_sub = p_sub.add_run("CAHIER DE FONCTIONNEMENT TOTAL &\nMANUEL D'UTILISATION EXHAUSTIF")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(20)
    run_sub.bold = True
    run_sub.font.color.rgb = SECONDARY_RGB

    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_before = Pt(12)
    p_desc.paragraph_format.space_after = Pt(30)
    p_desc.paragraph_format.line_spacing = 1.25
    run_desc = p_desc.add_run(
        "Guide pratique complet décrivant chaque module, chaque règle financière, "
        "chaque bouton et chaque flux de travail jusqu'au moindre détail.\n"
        "Comprend des exemples concrets en FCFA, des études de cas pas-à-pas, "
        "la gestion de caisse, des stocks, des crédits clients, des lots de marchandises "
        "et le pilotage par Copilote IA."
    )
    run_desc.font.size = Pt(11)
    run_desc.font.color.rgb = MUTED_RGB

    add_callout(
        "Ce manuel a été rédigé avec une volonté de clarté absolue : aucun terme technique "
        "n'est laissé sans explication simple. Que vous soyez caissier débutant, commerçant de quartier, "
        "gérant de magasin, grossiste ou importateur, ce guide vous montre précisément comment "
        "sécuriser chaque franc, stopper les fuites d'argent et faire fructifier votre entreprise.",
        title="ENGAGEMENT PÉDAGOGIQUE",
        bg_hex="ECFDF5",
        border_hex=PRIMARY
    )

    doc.add_page_break()

    # TABLE DES MATIERES
    add_h1("TABLE DES MATIÈRES")
    toc_items = [
        ("CHAPITRE 1", "Introduction, Vision et Fondamentaux Financiers"),
        ("CHAPITRE 2", "Inscription, Configuration Initiale (Onboarding en 5 Étapes)"),
        ("CHAPITRE 3", "Le Tableau de Bord (Dashboard) : Le Cockpit Décisionnel"),
        ("CHAPITRE 4", "Le Module Ventes & Point de Vente au Comptoir"),
        ("CHAPITRE 5", "La Caisse Journalière : Contrôle Physique & Clôture Rigoureuse"),
        ("CHAPITRE 6", "Gestion des Dépenses & Maîtrise des Charges d'Exploitation"),
        ("CHAPITRE 7", "Gestion des Produits, des Stocks et Réapprovisionnements"),
        ("CHAPITRE 8", "Carnet de Crédits Clients & Recouvrement des Impayés"),
        ("CHAPITRE 9", "Espace Pro : Gestion des Lots Commerciaux (Ballots, Friperie, Conteneurs)"),
        ("CHAPITRE 10", "Espace Pro : Croissance Marketing, Canaux Publicitaires & CAC"),
        ("CHAPITRE 11", "Espace Pro : Trésorerie Prévisionnelle & Anticipation des Pénuries"),
        ("CHAPITRE 12", "Espace Pro : Gestion des Livraisons, Colis et Paiements à la Livraison"),
        ("CHAPITRE 13", "Le Copilote IA COMMERCIA : Votre Conseiller Financier 24h/24"),
        ("CHAPITRE 14", "Architecture Multi-Boutiques, Gestion d'Équipe, Sécurité & Abonnements"),
        ("ANNEXE", "Glossaire du Commerçant Moderne : Comprendre les Termes de A à Z"),
    ]
    for ch_code, ch_title in toc_items:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(2)
        p_t.paragraph_format.space_after = Pt(4)
        run_c = p_t.add_run(f"{ch_code} : ")
        run_c.bold = True
        run_c.font.color.rgb = PRIMARY_RGB
        run_t = p_t.add_run(ch_title)
        run_t.font.color.rgb = SECONDARY_RGB

    doc.add_page_break()

    # CHAPITRE 1
    add_h1("CHAPITRE 1 : INTRODUCTION, VISION ET FONDAMENTAUX FINANCIERS")
    add_h2("1.1 Qu'est-ce que COMMERCIA ?")
    add_p(
        "COMMERCIA est une application web et mobile de gestion commerciale et financière "
        "conçue spécifiquement pour répondre aux réalités économiques des commerçants, boutiques, "
        "prestataires et entrepreneurs du continent africain. "
        "Elle fonctionne avec les devises locales (FCFA XAF, FCFA XOF, Franc Guinéen GNF, "
        "Franc Congolais CDF, Euro € et Dollar $) et intègre nativement les moyens de paiement "
        "du quotidien : Espèces (Cash direct) et Mobile Money (Orange Money, MTN MoMo, Wave, Moov)."
    )
    add_h2("1.2 Le Piège Mortel du Commerce Traditionnel : Chiffre d'Affaires vs Bénéfice Réel")
    add_p(
        "La cause n°1 de fermeture des commerces après seulement 1 ou 2 ans d'existence "
        "est la confusion mortelle entre le Chiffre d'Affaires (l'argent encaissé) "
        "et le Bénéfice Net (ce qui reste réellement dans la poche du commerçant)."
    )
    add_bullet("Le Chiffre d'Affaires (CA) : ", "C'est la totalité des sommes perçues par les ventes. Cet argent N'APPARTIENT PAS au commerçant. Une immense partie doit servir à payer les fournisseurs, le loyer, l'électricité, les emballages et les impôts.")
    add_bullet("Le Coût de Revient (Prix d'achat + Frais annexes) : ", "C'est l'argent obligatoire pour racheter la marchandise vendue. Si vous dépensez cet argent pour votre vie personnelle, vos rayons se vident irrémédiablement.")
    add_bullet("Le Bénéfice Net (La Marge Nette) : ", "C'est la seule et unique part qui constitue votre véritable gain personnel, calculée après avoir payé TOUTES les charges.")

    add_callout(
        "Moussa tient une boutique d'alimentation. Aujourd'hui, il a encaissé 100 000 FCFA dans son tiroir.\n"
        "Heureux de voir 100 000 FCFA en espèces, Moussa retire 30 000 FCFA pour faire plaisir à sa famille.\n"
        "QUE S'EST-IL RÉELLEMENT PASSÉ ?\n"
        "• Les marchandises vendues avaient coûté 85 000 FCFA à Moussa chez son grossiste.\n"
        "• La part proportionnelle du loyer et de l'électricité de la journée est de 5 000 FCFA.\n"
        "• Le vrai bénéfice de Moussa n'était que de : 100 000 - 85 000 - 5 000 = 10 000 FCFA !\n"
        "En prenant 30 000 FCFA, Moussa a en réalité mangé 20 000 FCFA du capital fournisseur.\n"
        "COMMERCIA empêche cette erreur fatale en calculant pour vous votre vraie marge à la seconde près !",
        title="ÉTUDE DE CAS : L'ILLUSION DANGEREUSE DE LA CAISSE PLEINE"
    )

    add_h2("1.3 Les 3 Modes de Gestion Proposés par COMMERCIA")
    add_p("L'application s'adapte au niveau d'organisation de chaque entreprise :")
    modes_data = [
        ["Mode A : Gestion Simplifiée", "Ventes + Dépenses + Marge estimée", "Ultra-rapide, sans inventaire lourd. Idéal pour salons de coiffure, snacks, services."],
        ["Mode B : Gestion avec Stock", "Articles + Coûts d'achat + Stocks + Alertes", "Recommandé pour boutiques, quincailleries, épiceries, prêt-à-porter."],
        ["Mode C : Gestion Complète", "Ventes + Caisse + Stocks + Carnet de Crédits + Rapports", "Le cockpit intégral 360° pour entrepreneurs organisés et PME avec vente à terme."]
    ]
    add_styled_table(["Mode de Gestion", "Fonctionnalités Incluses", "Public Recommandé"], modes_data, [2.0, 2.3, 2.5])

    # CHAPITRE 2
    add_h1("CHAPITRE 2 : CONFIGURATION INITIALE & ONBOARDING EN 5 ÉTAPES")
    add_p(
        "Lors de la première ouverture de COMMERCIA, un assistant interactif pas-à-pas "
        "configure votre environnement en 5 étapes simples :"
    )
    add_h3("Étape 1 : Identité de l'Entreprise")
    add_bullet("Nom commercial de la boutique ou de l'activité.")
    add_bullet("Devise principale : Par défaut FCFA (CEMAC XAF ou UEMOA XOF), GNF, CDF, EUR ou USD.")
    add_bullet("Numéro de téléphone et ville d'exploitation.")
    add_h3("Étape 2 : Secteur d'Activité")
    add_p(
        "COMMERCIA personnalise vos catégories de charges selon votre métier : "
        "Boutique/Supérette, Mode & Ballots friperie, Vente en ligne / WhatsApp, Restaurant / Bar, "
        "Salon de coiffure & Beauté, Atelier & Quincaillerie, ou Services / Agence."
    )
    add_h3("Étape 3 : Mode de Gestion (A, B ou C)")
    add_p("Activez ou désactivez en 1 clic les modules Stock et Carnet de crédits selon votre préférence.")
    add_h3("Étape 4 : Volume de Ventes Journalier Estimé")
    add_p(
        "Choisissez votre tranche de chiffre d'affaires prévisionnel (<25 000 FCFA, 25 000 à 100 000 FCFA, "
        "100 000 à 500 000 FCFA, ou >500 000 FCFA/jour). Cela permet au système de calibrer vos seuils "
        "d'alerte de dépenses anormales."
    )
    add_h3("Étape 5 : Formule d'Abonnement")
    add_bullet("Gratuit à vie : ", "Jusqu'à 15 ventes par jour, saisie de caisse et dépenses incluses.")
    add_bullet("Plan Pro (4 500 FCFA/mois) : ", "Ventes illimitées, stocks en temps réel, carnet de crédits, lots commerciaux, alertes WhatsApp et Copilote IA.")
    add_bullet("Plan Business (12 000 FCFA/mois) : ", "Multi-boutiques (jusqu'à 5 magasins), 10 comptes caissiers/gérants avec droits cloisonnés.")

    # CHAPITRE 3
    add_h1("CHAPITRE 3 : LE TABLEAU DE BORD (DASHBOARD) — LE COCKPIT DÉCISIONNEL")
    add_p(
        "Le Tableau de Bord est la première page qui s'affiche à la connexion. "
        "Il résume la santé financière globale de l'entreprise en temps réel."
    )
    add_h2("3.1 Les 5 Indicateurs Cardinaux (KPIs)")
    kpis = [
        ["1. Chiffre d'Affaires du Jour", "Total des ventes réalisées aujourd'hui (espèces, mobile money et crédits accordés). Permet de savoir si la journée a été dynamique."],
        ["2. Marge Nette Réelle", "Le véritable bénéfice net généré après déduction du coût des marchandises vendues et des dépenses enregistrées."],
        ["3. Dépenses Engagées", "Total des sorties d'argent opérées sur la période (achats de fournitures, loyer, factures, retraits)."],
        ["4. Trésorerie Disponible", "L'argent liquide immédiatement disponible en caisse et sur vos comptes Mobile Money."],
        ["5. Créances Clients en Souffrance", "L'argent de vos ventes qui dort dehors chez vos clients à crédit et qui doit être recouvré d'urgence."]
    ]
    add_styled_table(["Indicateur", "Signification & Rôle Décisionnel"], kpis, [2.3, 4.5])

    add_h2("3.2 Le Bandeau d'Alertes Proactives")
    add_bullet("Alerte Rouge Rupture : ", "Articles dont la quantité en stock est tombée sous le seuil d'alerte.")
    add_bullet("Alerte Orange Impayés : ", "Clients dont la date limite de remboursement de dette est dépassée.")
    add_bullet("Alerte Jaune Caisse : ", "Rappel que la caisse de la journée n'a pas encore été ouverte.")

    # CHAPITRE 4
    add_h1("CHAPITRE 4 : LE MODULE VENTES & ENCAISSEMENT AU COMPTOIR")
    add_p(
        "Le module Ventes permet d'enregistrer une transaction en moins de 15 secondes, "
        "sur téléphone, tablette ou ordinateur."
    )
    add_h2("4.1 Déroulement d'une Vente Pas-à-Pas")
    add_bullet("Étape 1 : Sélection des articles : ", "Recherchez le produit dans le catalogue ou utilisez le bouton 'Vente Libre'.")
    add_bullet("Étape 2 : Quantité et Rabais : ", "Ajustez la quantité et appliquez si besoin une remise commerciale en FCFA.")
    add_bullet("Étape 3 : Attribution Client (Optionnel) : ", "Associez la vente à un client de votre carnet (obligatoire si la vente est à crédit).")
    add_bullet("Étape 4 : Choix du Mode d'Encaissement : ", "Espèces (Cash), Mobile Money (Orange Money/MTN MoMo/Wave), Virement Bancaire, ou Vente à Crédit.")
    add_bullet("Étape 5 : Validation & Reçu : ", "Le stock est déduit, la caisse créditée et le ticket de caisse peut être imprimé ou envoyé sur WhatsApp.")

    add_callout(
        "Vous vendez 2 paires de chaussures de sport à 15 000 FCFA l'unité (Total 30 000 FCFA).\n"
        "• Le client demande une réduction : Vous saisissez une remise de 2 000 FCFA. Nouveau total = 28 000 FCFA.\n"
        "• Le client dispose de 10 000 FCFA en cash et promet de payer les 18 000 FCFA restants vendredi prochain.\n"
        "ACTION COMMERCIA :\n"
        "1. Vous sélectionnez le client 'Paul Demba'.\n"
        "2. Dans 'Montant Encaissé', vous tapez 10 000 FCFA en espèces.\n"
        "3. L'application calcule immédiatement un reste dû de 18 000 FCFA.\n"
        "4. En validant, 10 000 FCFA entrent automatiquement dans votre caisse journalière, "
        "2 paires sont déduites de votre stock, et une créance de 18 000 FCFA est créée sur la fiche de Paul Demba !",
        title="EXEMPLE COMPLET : VENTE AVEC REMISE, ACOMPTE ET CRÉDIT"
    )

    # CHAPITRE 5
    add_h1("CHAPITRE 5 : LA CAISSE JOURNALIÈRE — CONTRÔLE PHYSIQUE & CLÔTURE")
    add_p(
        "La caisse journalière est le garde-fou qui empêche les disparitions d'argent. "
        "Elle transforme un tiroir-caisse vulnérable en un système de contrôle comptable rigoureux."
    )
    add_h2("5.1 Le Cycle Quotidien de la Caisse")
    caisse_cycle = [
        ["1. Le Matin (Ouverture)", "Le caissier compte la monnaie présente dans le tiroir et saisit le Fond Initial (ex: 10 000 FCFA). La session est déclarée 'OUVERTE'."],
        ["2. La Journée (Flux continus)", "Chaque vente en liquide ajoute automatiquement des entrées. Chaque dépense payée en cash déduit des sorties. Tout retrait personnel du gérant est saisi avec le bouton '- Sortie de fonds'."],
        ["3. Le Soir (Clôture & Contrôle)", "Le caissier compte physiquement chaque billet et pièce dans le tiroir et saisit le 'Montant réel compté'. L'application compare avec le 'Solde théorique'."]
    ]
    add_styled_table(["Moment", "Opérations Réalisées & Impact"], caisse_cycle, [2.0, 4.8])

    add_h2("5.2 La Formule Mathématique de la Caisse")
    add_p("Solde Théorique = Fond Initial + Ventes Cash + Dépôts - Dépenses Cash - Retraits", bold_prefix="Formule : ")

    add_h2("5.3 Traitement des Écarts de Caisse")
    add_bullet("Écart = 0 FCFA : ", "Caisse parfaite. Tous les flux ont été fidèlement enregistrés.")
    add_bullet("Écart Négatif (ex: -1 500 FCFA) : ", "Déficit de caisse. Il manque de l'argent physique (erreur de monnaie, vol ou dépense oubliée).")
    add_bullet("Écart Positif (ex: +2 000 FCFA) : ", "Excédent de caisse. Il y a plus d'argent que prévu (vente oubliée d'être enregistrée).")
    
    add_callout(
        "RÈGLE COMMERCIA ANTI-FRAUDE :\n"
        "Si l'écart entre le comptage physique et le montant théorique dépasse 100 FCFA, "
        "le système refuse la clôture tant qu'un motif explicatif n'a pas été renseigné par l'utilisateur "
        "(ex: 'Erreur de rendu sur client de 14h' ou 'Achat d'eau minérale 500 FCFA non saisi'). "
        "Ce motif est gravé de manière inaltérable dans les journaux d'audit de l'entreprise !",
        title="SÉCURITÉ ET TRAÇABILITÉ STRICTE"
    )

    # CHAPITRE 6
    add_h1("CHAPITRE 6 : GESTION DES DÉPENSES & CHARGES D'EXPLOITATION")
    add_p(
        "La rentabilité d'un commerce dépend de la maîtrise de ses coûts. "
        "Le module Dépenses surveille chaque centime dépensé."
    )
    add_h2("6.1 Catégories Standards")
    add_bullet("Marchandises : ", "Achat direct de stock auprès des grossistes.")
    add_bullet("Loyer & Charges fixes : ", "Bail commercial de la boutique, gardiennage.")
    add_bullet("Électricité & Eau : ", "Factures d'énergie, recharge de compteur prépayé.")
    add_bullet("Transport & Ravitaillement : ", "Déplacement au marché, taxi-moto, carburant pour livraison.")
    add_bullet("Emballages & Packaging : ", "Sacs en plastique, cartons, ruban adhésif, étiquettes.")
    add_bullet("Salaires & Primes : ", "Rémunération du personnel et commissions sur vente.")
    add_bullet("Marketing & Publicité : ", "Campagnes Facebook Ads, affiches, influenceurs.")
    add_bullet("Divers / Imprévus : ", "Petites réparations, entretien, eau minérale.")

    # CHAPITRE 7
    add_h1("CHAPITRE 7 : GESTION DES ARTICLES, DES STOCKS ET DES RÉAPPROVISIONNEMENTS")
    add_p("Le module Articles offre une visibilité totale sur vos étagères.")
    add_h2("7.1 Fiche Produit")
    add_bullet("Référence & Nom : ", "Identifiant unique et désignation claire.")
    add_bullet("Coût de revient : ", "Prix d'achat unitaire + frais annexes.")
    add_bullet("Prix de vente : ", "Prix standard pratiqué au comptoir.")
    add_bullet("Seuil d'alerte : ", "Niveau critique déclenchant un signal de réapprovisionnement.")
    add_h2("7.2 Le PMP (Prix Moyen Pondéré)")
    add_p("PMP = [(10 x 3 000) + (10 x 4 000)] / 20 = 70 000 / 20 = 3 500 FCFA", bold_prefix="Exemple : ")

    # CHAPITRE 8
    add_h1("CHAPITRE 8 : CARNET DE CRÉDITS CLIENTS & RECOUVREMENT DES IMPAYÉS")
    add_p("Digitalisez votre carnet de crédits pour éliminer les impayés et sécuriser votre trésorerie.")
    add_h2("8.1 Fiche Créance")
    add_bullet("Total acheté, solde restant dû, date d'échéance et historique des règlements.")
    add_h2("8.2 Relance WhatsApp en 1 Clic")
    add_p("Générez un message cordial et pré-rédigé en 1 clic pour rappeler la date limite au client.")
    add_callout(
        "« Bonjour M. Ibrahim, nous vous rappelons cordialement que votre solde chez Espace Chic "
        "est de 15 000 FCFA (échéance au 25 Octobre). Merci de bien vouloir passer régler ou faire "
        "un dépôt Mobile Money au 690XXXXXX. Excellente journée ! »",
        title="EXEMPLE DE RELANCE AUTOMATIQUE"
    )

    # CHAPITRE 9
    add_h1("CHAPITRE 9 : MODULE PRO — GESTION DES LOTS COMMERCIAUX")
    add_p(
        "Module spécialisé pour importateurs, grossistes et friperie qui achètent des ballots ou conteneurs."
    )
    add_h2("9.1 Coût Global Débarqué (Landed Cost)")
    add_bullet("Prix d'achat initial, fret, douane, manutention, transport et pressing.")
    add_h2("9.2 Seuil de Rentabilité (Point Mort)")
    add_p("Dès que les ventes cumulées couvrent le coût total du lot, chaque vente suivante est du pur bénéfice net.")

    # CHAPITRE 10
    add_h1("CHAPITRE 10 : MODULE PRO — CROISSANCE, CANAUX & CAC")
    add_p("CAC = Dépenses Publicitaires du Canal / Nombre de Nouveaux Clients Acquis", bold_prefix="Formule du CAC : ")
    canaux_data = [
        ["Facebook Ads", "30 000 FCFA investis", "10 clients acquis", "3 000 FCFA / client", "Rentable si panier > 15 000 FCFA"],
        ["Statuts WhatsApp", "5 000 FCFA (connexion)", "25 clients réguliers", "200 FCFA / client", "EXCELLENT ROI (Priorité n°1)"],
        ["Flyers papier", "25 000 FCFA (impression)", "2 clients venus", "12 500 FCFA / client", "NON RENTABLE (À stopper)"]
    ]
    add_styled_table(["Canal Marketing", "Budget Dépensé", "Nouveaux Clients", "CAC Réel", "Décision"], canaux_data, [1.5, 1.3, 1.3, 1.3, 1.4])

    # CHAPITRE 11
    add_h1("CHAPITRE 11 : MODULE PRO — TRÉSORERIE PRÉVISIONNELLE & SCÉNARIOS")
    add_p("Projection de solde à 30, 60 et 90 jours pour anticiper les tensions de liquidités et sécuriser les échéances.")

    # CHAPITRE 12
    add_h1("CHAPITRE 12 : MODULE PRO — GESTION DES LIVRAISONS & COLIS")
    add_p("Attribution aux coursiers, suivi en temps réel et sécurisation du Paiement à la Livraison (Cash on Delivery).")

    # CHAPITRE 13
    add_h1("CHAPITRE 13 : LE COPILOTE IA COMMERCIA — VOTRE ANALYSTE 24H/24")
    add_p("Assistant intelligent qui détecte les produits dormants, les marges trop faibles et répond à vos questions en français.")

    # CHAPITRE 14
    add_h1("CHAPITRE 14 : ARCHITECTURE MULTI-BOUTIQUES, ÉQUIPE & ABONNEMENTS")
    add_p("Rôles cloisonnés (Propriétaire, Gérant, Caissier) et journal d'audit inaltérable. Paiement des abonnements par Mobile Money.")

    # ANNEXE
    add_h1("ANNEXE : GLOSSAIRE DU COMMERÇANT MODERNE")
    glossaire = [
        ["Chiffre d'Affaires (CA)", "Montant total de toutes les ventes réalisées, sans déduire aucune dépense ni achat."],
        ["Marge Brute", "Chiffre d'Affaires moins le coût d'achat des marchandises vendues."],
        ["Marge Nette (Bénéfice Net)", "Ce qui reste réellement après avoir payé le coût d'achat ET toutes les dépenses de fonctionnement."],
        ["Fond de Caisse", "Petite monnaie placée dans le tiroir au début de la journée pour pouvoir rendre la monnaie aux clients."],
        ["Écart de Caisse", "Différence constatée au soir entre l'argent compté dans le tiroir et ce que le système attendait."],
        ["Point Mort (Seuil de rentabilité)", "Le niveau de ventes exact à partir duquel l'activité couvre tous ses frais et commence à gagner de l'argent."],
        ["Créance Client", "Somme due par un client qui a acheté à crédit et qui doit vous rembourser plus tard."],
        ["PMP (Prix Moyen Pondéré)", "Méthode de calcul officielle de la valeur moyenne de votre stock lors d'achats successifs à des prix différents."]
    ]
    add_styled_table(["Terme Financier", "Définition Simple & Claire pour le Commerçant"], glossaire, [2.2, 4.6])

    doc.save(output_path)
    print(f"Document Word sauvegarde : {output_path}")

# =========================================================================
# 2. GÉNÉRATION PDF (.PDF) HAUTE DÉFINITION
# =========================================================================

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_number(self, page_count):
        if self._pageNumber > 1:
            self.saveState()
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor("#64748B"))
            # Header
            self.drawString(40, 800, "COMMERCIA — Cahier de Fonctionnement Intégral & Manuel d'Utilisation")
            self.setStrokeColor(colors.HexColor("#CBD5E1"))
            self.setLineWidth(0.5)
            self.line(40, 792, 555, 792)
            # Footer
            page_text = f"Page {self._pageNumber} sur {page_count}"
            self.drawRightString(555, 30, page_text)
            self.drawString(40, 30, "Document Confidentiel & Guide Opérationnel — Tous droits réservés")
            self.line(40, 42, 555, 42)
            self.restoreState()

def build_pdf_manual(pdf_path):
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=55,
        bottomMargin=55
    )
    
    styles = getSampleStyleSheet()
    c_primary = colors.HexColor("#059669")
    c_secondary = colors.HexColor("#1E293B")
    c_accent = colors.HexColor("#2563EB")
    c_muted = colors.HexColor("#64748B")
    c_callout_bg = colors.HexColor("#F0FDF4")
    c_callout_border = colors.HexColor("#059669")
    
    style_cover_pre = ParagraphStyle('CoverPre', fontName='Helvetica-Bold', fontSize=9, leading=12, textColor=c_muted, spaceAfter=15)
    style_cover_title = ParagraphStyle('CoverTitle', fontName='Helvetica-Bold', fontSize=32, leading=38, textColor=c_primary, spaceAfter=8)
    style_cover_sub = ParagraphStyle('CoverSub', fontName='Helvetica-Bold', fontSize=18, leading=24, textColor=c_secondary, spaceAfter=15)
    style_h1 = ParagraphStyle('PdfH1', fontName='Helvetica-Bold', fontSize=14, leading=18, textColor=c_primary, spaceBefore=18, spaceAfter=6, keepWithNext=True)
    style_h2 = ParagraphStyle('PdfH2', fontName='Helvetica-Bold', fontSize=11, leading=15, textColor=c_secondary, spaceBefore=12, spaceAfter=4, keepWithNext=True)
    style_body = ParagraphStyle('PdfBody', fontName='Helvetica', fontSize=9, leading=13.5, textColor=c_secondary, spaceAfter=5)
    style_bullet = ParagraphStyle('PdfBullet', parent=style_body, leftIndent=12, spaceAfter=3)
    style_callout = ParagraphStyle('PdfCallout', fontName='Helvetica', fontSize=8.5, leading=12.5, textColor=colors.HexColor("#334155"))
    style_th = ParagraphStyle('PdfTH', fontName='Helvetica-Bold', fontSize=8.5, leading=11, textColor=colors.white)
    style_td = ParagraphStyle('PdfTD', fontName='Helvetica', fontSize=8, leading=11, textColor=c_secondary)

    def make_callout(title, text):
        data = [[Paragraph(f"<b>📌 {title}</b><br/>{text}", style_callout)]]
        t = Table(data, colWidths=[515])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), c_callout_bg),
            ('BOX', (0,0), (-1,-1), 1.5, c_callout_border),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('LEFTPADDING', (0,0), (-1,-1), 10),
            ('RIGHTPADDING', (0,0), (-1,-1), 10),
        ]))
        return t

    def make_table(headers, rows, widths):
        data = [[Paragraph(h, style_th) for h in headers]]
        for r in rows:
            data.append([Paragraph(str(c), style_td) for c in r])
        t = Table(data, colWidths=widths)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), c_primary),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('LEFTPADDING', (0,0), (-1,-1), 6),
            ('RIGHTPADDING', (0,0), (-1,-1), 6),
        ]))
        return t

    story = []

    # COUVERTURE
    story.append(Spacer(1, 40))
    story.append(Paragraph("RÉPUBLIQUE DU COMMERCE & GESTION MODERNE<br/>MANUEL OFFICIEL D'EXPLOITATION OPÉRATIONNELLE", style_cover_pre))
    story.append(Paragraph("COMMERCIA", style_cover_title))
    story.append(Paragraph("CAHIER DE FONCTIONNEMENT TOTAL &<br/>MANUEL D'UTILISATION EXHAUSTIF", style_cover_sub))
    story.append(HRFlowable(width="100%", thickness=2.5, color=c_primary, spaceBefore=2, spaceAfter=15))
    story.append(Paragraph(
        "Ce manuel détaille chaque écran, règle comptable, bouton et flux de travail de l'application "
        "COMMERCIA. Conçu avec une clarté totale pour que chaque commerçant, caissier, gérant ou importateur "
        "comprenne immédiatement comment sécuriser son argent, piloter sa caisse, maîtriser ses stocks et marges "
        "et développer sereinement son entreprise.",
        style_body
    ))
    story.append(Spacer(1, 10))
    story.append(make_callout(
        "ENGAGEMENT DE TRANSPARENCE ET DE SIMPLICITÉ",
        "Aucun jargon abstrait : chaque notion financière est illustrée d'exemples de la vie courante "
        "en FCFA (boutique d'alimentation, prêt-à-porter, salon, restaurant, importateur de friperie). "
        "En suivant ce guide, vous éliminerez les fuites de caisse dès le premier jour !"
    ))
    story.append(PageBreak())

    # TABLE DES MATIERES
    story.append(Paragraph("TABLE DES MATIÈRES GÉNÉRALE", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=c_primary, spaceBefore=2, spaceAfter=10))
    tocs = [
        "1. Introduction, Vision & Fondamentaux Financiers (Ne plus confondre CA et Marge)",
        "2. Configuration Initiale & Onboarding en 5 Étapes",
        "3. Le Tableau de Bord (Dashboard) : Le Cockpit Décisionnel",
        "4. Le Module Ventes & Point de Vente au Comptoir",
        "5. La Caisse Journalière : Contrôle Physique & Clôture Rigoureuse",
        "6. Gestion des Dépenses & Maîtrise des Charges d'Exploitation",
        "7. Gestion des Articles, Stocks & Calcul du PMP",
        "8. Carnet de Crédits Clients & Relances WhatsApp en 1 Clic",
        "9. Module Pro : Gestion des Lots Commerciaux (Ballots, Friperie, Conteneurs)",
        "10. Module Pro : Croissance Marketing, Canaux Publicitaires & Calcul du CAC",
        "11. Module Pro : Trésorerie Prévisionnelle & Anticipation des Pénuries",
        "12. Module Pro : Gestion des Livraisons, Colis et Paiements à la Livraison",
        "13. Le Copilote IA COMMERCIA : Votre Analyste Financier Personnel 24h/24",
        "14. Multi-Boutiques, Équipe, Audit Logs & Modalités d'Abonnement",
        "ANNEXE : Glossaire Complet du Commerçant Moderne (De A à Z)"
    ]
    for idx, item in enumerate(tocs, 1):
        story.append(Paragraph(f"<b>Chapitre {idx}</b> : {item}" if idx <= 14 else f"<b>{item}</b>", style_body))
    story.append(PageBreak())

    # CONTENU COMPLET DES CHAPITRES
    # Chapitre 1
    story.append(Paragraph("CHAPITRE 1 : INTRODUCTION, VISION & FONDAMENTAUX", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=c_primary, spaceBefore=2, spaceAfter=8))
    story.append(Paragraph("1.1 Pourquoi COMMERCIA ?", style_h2))
    story.append(Paragraph("COMMERCIA répond au drame silencieux des commerçants : travailler 14 heures par jour, voir des centaines de milliers de FCFA entrer dans le tiroir, mais se retrouver en faillite sans comprendre où est parti l'argent.", style_body))
    story.append(Paragraph("1.2 L'erreur mortelle : Chiffre d'Affaires vs Marge Nette", style_h2))
    story.append(Paragraph("Le Chiffre d'Affaires est l'argent brut qui passe dans vos mains. Il ne vous appartient pas : 80% à 90% doit rembourser vos fournisseurs et payer les factures. La Marge Nette est votre véritable gain.", style_body))
    story.append(make_callout(
        "ÉTUDE DE CAS : L'ILLUSION DANGEREUSE DE LA CAISSE PLEINE",
        "Moussa encaisse 100 000 FCFA aujourd'hui. Il prend 30 000 FCFA pour des dépenses familiales.<br/>"
        "RÉALITÉ : La marchandise vendue lui a coûté 85 000 FCFA chez le grossiste. Le loyer du jour est de 5 000 FCFA.<br/>"
        "Son vrai bénéfice était de 10 000 FCFA ! En retirant 30 000 FCFA, Moussa a mangé 20 000 FCFA du capital qui devait racheter du stock. C'est ainsi que les rayons se vident !"
    ))
    story.append(Spacer(1, 8))

    # Chapitre 2 & 3
    story.append(Paragraph("CHAPITRE 2 : CONFIGURATION INITIALE EN 5 ÉTAPES", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=c_primary, spaceBefore=2, spaceAfter=8))
    story.append(Paragraph("1. Identité (Nom, Devise FCFA/GNF/CDF/EUR/USD) | 2. Secteur métier | 3. Mode A (Simple), B (Stock) ou C (Complet) | 4. Volume journalier | 5. Choix de formule (Gratuit 15 ventes/j, Pro 4 500 FCFA, Business 12 000 FCFA).", style_body))

    story.append(Paragraph("CHAPITRE 3 : LE TABLEAU DE BORD (DASHBOARD)", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=c_primary, spaceBefore=2, spaceAfter=8))
    story.append(Paragraph("Le cockpit affiche 5 KPIs cardinaux : CA du jour, Marge nette réelle calculée, Dépenses engagées, Trésorerie disponible et Créances clients dehors avec alertes immédiates.", style_body))

    # Chapitre 4 & 5
    story.append(Paragraph("CHAPITRE 4 : LE MODULE VENTES & COMPTOIR", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=c_primary, spaceBefore=2, spaceAfter=8))
    story.append(Paragraph("Encaissement ultra-rapide avec sélection d'articles ou vente libre. Prise en charge des remises, acomptes, crédits et moyens de paiement (Cash, Orange Money, MTN MoMo, Wave). Génération instantanée du ticket de caisse et partage WhatsApp au client.", style_body))
    story.append(make_callout(
        "EXEMPLE CONCRET DE VENTE MIXTE (CASH + CRÉDIT)",
        "Vente de 2 paires de chaussures à 15 000 FCFA (Total 30 000 FCFA) avec remise de 2 000 FCFA -> Total 28 000 FCFA.<br/>"
        "Le client paie 10 000 FCFA cash et doit 18 000 FCFA. COMMERCIA crédite 10 000 FCFA dans la caisse, déstocke 2 paires et crée une créance de 18 000 FCFA dans la fiche client !"
    ))

    story.append(Paragraph("CHAPITRE 5 : LA CAISSE JOURNALIÈRE — CONTRÔLE & CLÔTURE", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=c_primary, spaceBefore=2, spaceAfter=8))
    story.append(Paragraph("<b>Matin :</b> Déclaration du Fond Initial.<br/><b>Journée :</b> Enregistrement des flux et sorties manuelles du gérant.<br/><b>Soir :</b> Comptage physique des billets/pièces. L'application calcule l'écart. Si l'écart dépasse 100 FCFA, un motif d'explication obligatoire doit être saisi.", style_body))

    # Chapitres 6, 7 & 8
    story.append(Paragraph("CHAPITRE 6 : DÉPENSES & CHARGES D'EXPLOITATION", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=c_primary, spaceBefore=2, spaceAfter=8))
    story.append(Paragraph("Surveillance stricte des postes de coûts (Marchandise, Loyer, Électricité, Salaires, Transport, Packaging, Publicité) et distinction entre charges fixes et charges variables.", style_body))

    story.append(Paragraph("CHAPITRE 7 : ARTICLES, STOCKS & CALCUL DU PMP", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=c_primary, spaceBefore=2, spaceAfter=8))
    story.append(Paragraph("Gestion du catalogue avec Prix Moyen Pondéré (PMP) lors des achats successifs, sorties automatiques sur vente, ajustements de casse et alertes de seuil minimum.", style_body))

    story.append(Paragraph("CHAPITRE 8 : CRÉDITS CLIENTS & RELANCES WHATSAPP", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=c_primary, spaceBefore=2, spaceAfter=8))
    story.append(Paragraph("Fin des impayés grâce au suivi par échéance et au bouton de relance WhatsApp en 1 clic avec message personnalisé et courtois pré-rempli.", style_body))

    # Chapitres 9, 10, 11, 12, 13 & 14
    story.append(Paragraph("CHAPITRE 9 : MODULE PRO — GESTION DES LOTS COMMERCIAUX", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=c_primary, spaceBefore=2, spaceAfter=8))
    story.append(Paragraph("Idéal pour grossistes et ballots de friperie : Calcul du Coût Global Débarqué (Achat + Fret + Douane + Manutention + Pressing), répartition par pièce, décompte des rebuts et suivi du franchissement du Point Mort (Seuil de rentabilité).", style_body))

    story.append(Paragraph("CHAPITRE 10 : MODULE PRO — CROISSANCE, CANAUX & CAC", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=c_primary, spaceBefore=2, spaceAfter=8))
    story.append(Paragraph("Mesure du Coût d'Acquisition Client (CAC = Dépenses pub / Nouveaux clients) sur Facebook, WhatsApp et prospection terrain pour éliminer les publicités non rentables.", style_body))

    story.append(Paragraph("CHAPITRES 11 & 12 : TRÉSORERIE & LIVRAISONS", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=c_primary, spaceBefore=2, spaceAfter=8))
    story.append(Paragraph("Prévisionnel de trésorerie à 30/60/90 jours et encadrement strict du Paiement à la Livraison (Cash on Delivery) avec les livreurs.", style_body))

    story.append(Paragraph("CHAPITRES 13 & 14 : COPILOTE IA, SÉCURITÉ & MULTI-BOUTIQUES", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=c_primary, spaceBefore=2, spaceAfter=8))
    story.append(Paragraph("Conseiller IA 24h/24 analysant vos marges faibles et produits dormants. Gestion d'équipe avec rôles cloisonnés (Propriétaire, Gérant, Caissier) et journal d'audit inaltérable.", style_body))
    story.append(Spacer(1, 8))

    # ANNEXE GLOSSAIRE
    story.append(Paragraph("ANNEXE : GLOSSAIRE DU COMMERÇANT MODERNE", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=c_primary, spaceBefore=2, spaceAfter=8))
    glossaire_pdf = [
        ["Chiffre d'Affaires (CA)", "Montant total de toutes les ventes sans déduire aucune dépense."],
        ["Marge Nette (Bénéfice Net)", "Ce qui reste réellement après paiement de la marchandise et de TOUTES les charges."],
        ["Fond de Caisse", "Monnaie placée dans le tiroir le matin pour rendre la monnaie aux clients."],
        ["Écart de Caisse", "Différence au soir entre le comptage physique réel et le solde théorique."],
        ["Point Mort", "Seuil de ventes à partir duquel tous les coûts sont amortis et le profit commence."],
        ["PMP (Prix Moyen Pondéré)", "Formule officielle de calcul de la valeur moyenne de stock lors d'achats à prix variés."]
    ]
    story.append(make_table(["Terme", "Définition Simple"], glossaire_pdf, [150, 365]))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"Document PDF sauvegarde : {pdf_path}")

if __name__ == '__main__':
    base_dir = r"c:\Users\HACKER SUR MESURE\Downloads\APPLICATON DE GESTION COMMERCIAL COMMERCIA"
    docx_file = os.path.join(base_dir, "COMMERCIA_Cahier_de_Fonctionnement_Integral.docx")
    pdf_file = os.path.join(base_dir, "COMMERCIA_Cahier_de_Fonctionnement_Integral.pdf")
    
    print("Generation Word...")
    build_docx_manual(docx_file)
    print("Generation PDF...")
    build_pdf_manual(pdf_file)
    print("SUCCES TOTAL !")
